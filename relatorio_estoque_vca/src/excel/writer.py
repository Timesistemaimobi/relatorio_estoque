from __future__ import annotations

from copy import copy
from typing import Dict, Tuple

from openpyxl import Workbook
from openpyxl.cell import Cell

from .styles import copy_style


def _apply_zebra(
    cell_left: Cell,
    cell_right: Cell,
    styles: Dict[str, Cell],
    zebra_index: int,
    bold: bool = False,
) -> None:
    style_key = "zebra_1" if zebra_index % 2 == 0 else "zebra_2"
    source_left = styles[f"{style_key}_left"]
    source_right = styles[f"{style_key}_right"]
    copy_style(cell_left, source_left)
    copy_style(cell_right, source_right)
    if bold:
        left_font = copy(cell_left.font)
        left_font.bold = True
        cell_left.font = left_font
        right_font = copy(cell_right.font)
        right_font.bold = True
        cell_right.font = right_font


def write_report(
    output_path: str,
    template_styles: Dict[str, Cell],
    col_widths: Dict[str, float],
    aggregates_by_emp: Dict[str, dict],
    avg_prices_by_emp: Dict[str, Dict[str, float]],
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for empreendimento, aggregates in aggregates_by_emp.items():
        ws = wb.create_sheet(title=str(empreendimento)[:31])
        for col, width in col_widths.items():
            if width:
                ws.column_dimensions[col].width = width

        ws.merge_cells("B2:C2")
        ws["B2"].value = empreendimento
        copy_style(ws["B2"], template_styles["empreendimento_title"])

        ws.merge_cells("E2:F2")
        ws["E2"].value = f"VALOR MÉDIO - {empreendimento}"
        copy_style(ws["E2"], template_styles["avg_title"])

        ws["E3"].value = "Tipologia"
        ws["F3"].value = "Valor Médio"
        copy_style(ws["E3"], template_styles["avg_header_left"])
        copy_style(ws["F3"], template_styles["avg_header_right"])

        avg_prices = avg_prices_by_emp.get(empreendimento, {})
        avg_row = 4
        for tipologia in sorted(avg_prices):
            ws[f"E{avg_row}"].value = tipologia
            ws[f"F{avg_row}"].value = avg_prices[tipologia]
            copy_style(ws[f"E{avg_row}"], template_styles["avg_row_left"])
            copy_style(ws[f"F{avg_row}"], template_styles["avg_row_right"])
            avg_row += 1

        current_row = 4
        for etapa, data in aggregates.items():
            ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
            ws.cell(row=current_row, column=2, value=etapa)
            copy_style(ws.cell(row=current_row, column=2), template_styles["etapa_title"])
            current_row += 1

            ws.cell(row=current_row, column=2, value="Situação / Tipologia")
            ws.cell(row=current_row, column=3, value="Quantidade")
            copy_style(ws.cell(row=current_row, column=2), template_styles["header_left"])
            copy_style(ws.cell(row=current_row, column=3), template_styles["header_right"])
            current_row += 1

            zebra_index = 0
            for motivo, total, tipologias in data["bloqueios"]:
                left = ws.cell(row=current_row, column=2, value=f"BLOQUEADA - {motivo}")
                right = ws.cell(row=current_row, column=3, value=total)
                _apply_zebra(left, right, template_styles, zebra_index, bold=True)
                current_row += 1
                zebra_index += 1
                for tipologia, qtd in tipologias:
                    left = ws.cell(row=current_row, column=2, value=f"  {tipologia}")
                    right = ws.cell(row=current_row, column=3, value=qtd)
                    _apply_zebra(left, right, template_styles, zebra_index)
                    current_row += 1
                    zebra_index += 1

            for status, total, tipologias in data["status"]:
                left = ws.cell(row=current_row, column=2, value=status)
                right = ws.cell(row=current_row, column=3, value=total)
                _apply_zebra(left, right, template_styles, zebra_index, bold=True)
                current_row += 1
                zebra_index += 1
                for tipologia, qtd in tipologias:
                    left = ws.cell(row=current_row, column=2, value=f"  {tipologia}")
                    right = ws.cell(row=current_row, column=3, value=qtd)
                    _apply_zebra(left, right, template_styles, zebra_index)
                    current_row += 1
                    zebra_index += 1

            left = ws.cell(row=current_row, column=2, value="TOTAL GERAL")
            right = ws.cell(row=current_row, column=3, value=data["total"])
            copy_style(left, template_styles["total_left"])
            copy_style(right, template_styles["total_right"])
            current_row += 2

        ws.freeze_panes = "B6"

    wb.save(output_path)
