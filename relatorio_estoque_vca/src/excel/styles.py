from __future__ import annotations

from copy import copy
from typing import Dict, Tuple

from openpyxl import load_workbook
from openpyxl.cell import Cell


def copy_style(target: Cell, source: Cell) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def load_styles(template_path: str) -> Tuple[Dict[str, Cell], Dict[str, float]]:
    wb = load_workbook(template_path)
    ws = wb.active
    styles = {
        "empreendimento_title": ws["B2"],
        "etapa_title": ws["B4"],
        "header_left": ws["B5"],
        "header_right": ws["C5"],
        "zebra_1_left": ws["B6"],
        "zebra_1_right": ws["C6"],
        "zebra_2_left": ws["B7"],
        "zebra_2_right": ws["C7"],
        "total_left": ws["B8"],
        "total_right": ws["C8"],
        "avg_title": ws["E2"],
        "avg_header_left": ws["E3"],
        "avg_header_right": ws["F3"],
        "avg_row_left": ws["E4"],
        "avg_row_right": ws["F4"],
    }
    col_widths = {}
    for col in ["A", "B", "C", "D", "E", "F"]:
        dimension = ws.column_dimensions[col]
        col_widths[col] = dimension.width
    return styles, col_widths
